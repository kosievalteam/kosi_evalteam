import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
d=pd.read_csv('all.csv')
mp={'산업통상자원부':'산업통상부','환경부':'기후에너지환경부','특허청':'지식재산처','방송통신위원회':'방송미디어통신위원회','여성가족부':'성평등가족부','통계청':'국가데이터처','문화재청':'국가유산청'}
d['소관(2026기준)']=d.소관명.replace(mp).str.replace(' ','')
base=['고용노동부','과학기술정보통신부','국토교통부','금융위원회','기후에너지환경부','농림축산식품부','문화체육관광부','방송미디어통신위원회','산림청','산업통상부','식품의약품안전처','조달청','중소벤처기업부','지식재산처','해양수산부','행정안전부']
sel={2024:set(base+['관세청']),2025:set(base),2026:set(base+['교육부','국가유산청','기상청','보건복지부','성평등가족부','우주항공청'])}
kw='중소|소상공인|소공인|창업|벤처|스타트업|중견|기업|수출|전통시장|상권|사업화|기술이전|판로|마케팅|바우처|융자|보증|일자리|고용'
excl='기본경비|인건비|운영비|청사|정보화|국제분담금|채무|상환|예탁|전출|전입|예치|보전'
dbn={2024:8290,2025:7939,2026:8587}
summ=[];bys=[];cands=[]
for y in [2024,2025,2026]:
    x=d[d.회계연도==y].copy()
    x['지원사업확인소관']=x['소관(2026기준)'].isin(sel[y]).map({True:'예',False:'아니오'})
    o=x[x.지원사업확인소관=='아니오']
    hit=o[(o.세부사업명.str.contains(kw)|o.단위사업명.str.contains(kw))&~(o.세부사업명.str.contains(excl)|o.단위사업명.str.contains(excl))]
    cands.append(hit)
    summ.append({'회계연도':y,'국회확정 세부사업(회계별 행)':len(x),'고유 세부사업(소관·사업명)':x[['소관명','세부사업명']].drop_duplicates().shape[0],
      '소관 수':x.소관명.nunique(),'지원사업 확인 소관 수':len(sel[y]),'지원사업 미확인 소관 수':o.소관명.nunique(),
      '기존 설명자료 DB 행':dbn[y],'설명자료 미확보 행(차이)':len(x)-dbn[y],
      '기본경비·인건비 행':int((x.단위사업명.str.contains('기본경비|인건비')|x.세부사업명.str.contains('기본경비|인건비')).sum()),
      '미확인 소관 내 기업지원 어휘 사업(후보)':len(hit)})
    g=x.groupby(['소관명','지원사업확인소관']).agg(세부사업수=('세부사업명','size'),예산_백만원=('예산액','sum')).reset_index()
    g['후보사업수']=g.소관명.map(hit.groupby('소관명').size()).fillna(0).astype(int); g.insert(0,'회계연도',y); bys.append(g)
out='/home/user/kosi_evalteam/planning/국회확정예산_전부처_검토범위_진단.xlsx'
c=pd.concat(cands)[['회계연도','소관명','회계명','분야명','부문명','프로그램명','단위사업명','세부사업명','예산액']].rename(columns={'예산액':'예산액(백만원)'})
notes=pd.DataFrame({'주석':['자료: 열린재정 재정사업 설명자료(분류별검색) 국회확정 예산 2024~2026, 정책평가팀 DB(all_biz, biz)',
 '지원사업 확인 소관: 정책평가팀 DB(biz, 중앙부처)에 해당 연도 지원사업이 1건 이상 있는 소관(2026년 조직 기준 명칭으로 대조)',
 '설명자료 미확보 행: 국회확정 예산 목록 행 수와 기존 사업설명자료 DB(all_biz) 행 수의 차이',
 '검토 후보: 지원사업 미확인 소관의 세부사업 중 단위사업·세부사업명에 기업 지원 관련 어휘가 있고 운영경비·채무 관련 어휘가 없는 사업. 지원사업 해당 여부는 판정 필요',
 '예산액 단위: 백만원(총계 기준, 내부거래·보전지출 포함)']})
sheets={'연도별 요약':pd.DataFrame(summ),'소관별 현황':pd.concat(bys).sort_values(['회계연도','지원사업확인소관','세부사업수'],ascending=[True,True,False]),'검토 후보(미확인 소관)':c,'주석':notes}
with pd.ExcelWriter(out,engine='openpyxl') as w:
    for n,df in sheets.items():
        df.to_excel(w,sheet_name=n,index=False)
        ws=w.sheets[n]
        for cell in ws[1]:
            cell.font=Font(bold=True); cell.fill=PatternFill('solid',fgColor='D9E2F3'); cell.alignment=Alignment(wrap_text=True,vertical='center')
        for col in ws.columns:
            L=max(len(str(v.value or '')) for v in col[:200]); ws.column_dimensions[col[0].column_letter].width=min(max(10,L*1.6),70)
        ws.freeze_panes='A2'
print(pd.DataFrame(summ).T.to_string())
